"""Use bundled python-docx; retain the supplied PART V design in local copies."""
from pathlib import Path
import hashlib
import json
import os
import re
import sys

from docx import Document

OUT = Path(__file__).resolve().parent


def lp(p):
    path = os.path.abspath(str(p))
    return '\\\\?\\' + path if os.name == 'nt' and not path.startswith('\\\\?\\') else path


def main():
    sources = [('PART_V_Path_B_Seven_Variable_Figure_Structure.docx', 'PART_V_Figure_Structure_Checked.docx'),
               ('Path_B_Elsevier_Figure_Captions.docx', 'Path_B_Figure_Captions_Checked.docx')]
    records = []
    for src, dst in sources:
        source = OUT / 'inputs' / src
        document = Document(lp(source))
        edits = []
        paragraphs = list(document.paragraphs) + [p for t in document.tables for row in t.rows
                                                for c in row.cells for p in c.paragraphs]
        for p in paragraphs:
            if 'independent Unconstrained reference' in p.text:
                old = p.text
                for run in p.runs:
                    run.text = run.text.replace('independent Unconstrained reference', 'common Unconstrained reference')
                edits.append(dict(before=old, after=p.text,
                                  reason='Unconstrained shares the frozen physical realizations; it is not an independent sample.'))
            if 'mapped with frozen Q1–Q4 SOVI hues;' in p.text:
                old = p.text
                for run in p.runs:
                    run.text = run.text.replace('mapped with frozen Q1–Q4 SOVI hues;',
                        'mapped with frozen Q1–Q4 SOVI hues and common burden-to-lightness scaling;')
                edits.append(dict(before=old, after=p.text,
                                  reason='Equal service-deficit magnitudes use the same luminance across quartile families.'))
            if 'Independent-realization service-loss contrasts, with a mean marker and saved realization range.' in p.text:
                old = p.text
                for run in p.runs:
                    run.text = run.text.replace(
                        'Independent-realization service-loss contrasts, with a mean marker and saved realization range.',
                        'The empirical distribution of independent-validation service-loss differences, with the mean marked.')
                assert p.text != old
                edits.append(dict(before=old, after=p.text,
                                  reason='Current S4C shows an empirical distribution and mean, not a separate realization-range whisker.'))
            if 'Removing source connectivity at fixed station threshold.' in p.text:
                old = p.text
                for run in p.runs:
                    run.text = run.text.replace('Removing source connectivity at fixed station threshold.',
                                                'Removing source connectivity at the fixed 0.50 station threshold.')
                edits.append(dict(before=old, after=p.text, reason='Make the frozen source-only control threshold explicit.'))
        all_text = '\n'.join(p.text for p in document.paragraphs)
        all_text += '\n'.join(p.text for t in document.tables for row in t.rows for c in row.cells for p in c.paragraphs)
        assert not any(x in all_text for x in ['2,291', 'S12', 'S13', 'five-class'])
        if 'Captions' in dst:
            main = re.findall(r'^Figure (\d+)\.', all_text, re.M)
            supplement = re.findall(r'^Supplementary Figure S(\d+)\.', all_text, re.M)
            assert main == [str(i) for i in range(1, 8)] and supplement == [str(i) for i in range(1, 11)]
            assert 'all seven P2 variables' in all_text and 'Six empirical feature distributions' in all_text
            assert '95% bootstrap confidence intervals for mean effects' in all_text
            assert 'service-loss bars are means without whiskers' in all_text
        table_checks = []
        for t in document.tables:
            available = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
            total = sum(c.width for c in t.columns)
            assert total <= available + 9144
            table_checks.append(dict(rows=len(t.rows), columns=len(t.columns), width_inches=total/914400,
                                     available_inches=available/914400, autofit=t.autofit))
        target = OUT / dst
        assert not target.exists() or '--revise' in sys.argv, target
        document.save(lp(target))
        records.append(dict(source=src, target=dst,
                            source_sha256=hashlib.sha256(open(lp(source), 'rb').read()).hexdigest(),
                            output_sha256=hashlib.sha256(open(lp(target), 'rb').read()).hexdigest(),
                            changes=edits, original_page_sizes_inches=[
                                [s.page_width.inches, s.page_height.inches] for s in document.sections],
                            table_checks=table_checks, source_edited=False, render_qa_pending=True))
    with open(lp(OUT / 'DOCX_CONTENT_AND_STYLE_QA.json'), 'w', encoding='utf-8') as f:
        json.dump(records, f, indent=2, ensure_ascii=True)
    print('Two local DOCX copies prepared, supplied styles preserved; source inputs unchanged.')


if __name__ == '__main__':
    main()
