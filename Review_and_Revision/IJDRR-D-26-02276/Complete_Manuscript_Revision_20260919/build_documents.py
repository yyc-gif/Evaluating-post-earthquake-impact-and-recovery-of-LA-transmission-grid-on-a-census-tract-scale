"""Editable manuscript package from frozen results; document work only."""
from pathlib import Path
import re, shutil, json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
SRC=ROOT.parent/'Targeted_Methodological_Strengthening_20260919'
SOURCE=ROOT/'source'; SOURCE.mkdir(exist_ok=True)
NAMES=['REVISED_MANUSCRIPT','SUPPLEMENTARY_MATERIAL','RESPONSE_TO_REVIEWERS','CLAIM_EVIDENCE_CROSSWALK']
for name in NAMES:
    target=SOURCE/(name+'.md.in')
    if not target.exists(): shutil.copyfile(ROOT/(name+'.md'),target)
for name in ['REVISED_MANUSCRIPT','SUPPLEMENTARY_MATERIAL']:
    p=SOURCE/(name+'.md.in'); t=p.read_text(encoding='utf-8')
    t=t.replace('population- and hospital-weighted priority sequences reduce','two weighted-priority sequences reduce')
    t=t.replace('less than 10⁻¹² h','less than 2×10⁻¹² h')
    p.write_text(t,encoding='utf-8')

D=pd.read_csv(SRC/'METRIC_DISTRIBUTIONS.csv')
E=pd.read_csv(SRC/'PAIRED_EFFECTS.csv')
W=pd.read_csv(SRC/'WINNER_LOSER_POPULATION.csv')
STR=['Hospital-first','GA-Balanced','GA-HospFirst','GA-Efficiency']
LAB=['Hospital-first','Balanced','HospFirst','Efficiency']
def mean(s,m,c='corrected_baseline'):
    return float(D[(D.condition==c)&(D.strategy==s)&(D.metric==m)].iloc[0]['mean'])
def effect(s,m,c='corrected_baseline',contrast='strategy_minus_HF',digits=3):
    r=E[(E.contrast==contrast)&(E.condition==c)&(E.strategy==s)&(E.metric==m)].iloc[0]
    return f'{r.mean_paired_difference:+.{digits}f} [{r.ci95_low:+.{digits}f}, {r.ci95_high:+.{digits}f}]'.replace('-', '−')
TABLES={}
def table(key,title,heads,rows,note='',widths=None):
    TABLES.setdefault(key,[]).append((title,heads,rows,note,widths))
table('TABLE1','Table 1. Study inputs, decisions, and interpretation boundaries.',
      ['Component','Adopted representation','Evidence or limitation'],[
      ['Task assets','302 assets in the primary component of a 310-record inventory','Mixed owners; all resolved service attachments fall in this component.'],
      ['Hazard and damage','Fixed CGS 2%-in-50-year PGA; voltage-class fragility; 32 stored damage vectors','Conditional scenario, not 32 earthquake ruptures; no new correlated intensity fields.'],
      ['Task duration','Positive-conditioned Normal DS1–DS4: (1,0.5), (6,3), (12,4), (36,12) h','Scenario on-site actions, not calibrated permanent repair times.'],
      ['Resources and roads','57 pooled crews; 11 origins; directed base–task and task–task travel','Allocation and access proxies; static roads; no crew-count contrast.'],
      ['Dispatch','Full sequence filtered by DS>0; earliest free crew; completion = arrival + duration','Duration affects crew release, never pending-task order.'],
      ['Raw restoration','DS residuals 1, 0.50, 0.09, 0.04, 0.03; damaged assets become 1 at completion','Completion-step assumption; no additional repair CDF.'],
      ['Network availability','Functionality threshold 0.5; components with an active reference source; fixed 21-source set','Connectivity proxy; no power-flow, capacity, or switching solution.'],
      ['Service candidates','817 strict-SCE tracts; 196 candidates: 113 A, 71 B, 12 C','Official candidate evidence; equal weights are not actual service shares.'],
      ['Population and missingness','3,572,152 total; 3,520,382 in 805 partly/fully represented tracts','12 fully unresolved tracts, 51,770 people, retain NA burden.'],
      ['Sequences','Hospital-first; deterministic Balanced/HospFirst initializers; retained Efficiency search candidate','Fixed ex-ante lists; no global-optimality or equity-optimization claim.'],
      ['Candidate-influence contrast','Relative B influence λ=0.5 or 2 in mixed A/B tracts','Fixed decisions; total resolved and C masses unchanged; offline burdens only.'],
      ['Relative-duration contrast','Stored DS4 durations ×2; all other inputs and sequences fixed','Paired rescheduling; no new draws or reoptimization.'],
      ],'Provider and scenario details appear in Sections 2.1–2.7 and Supplement S1–S4.',[1.15,2.55,3.1])
metrics=[('Population-normalized burden (h)','population_normalized_burden_hr'),
 ('Availability-consistent burden (h)','population_resolved_mass_burden_hr'),
 ('Population resolved T80 (h)','population_resolved_T80_hr'),
 ('Hospital-tract burden (h)','hospital_mean_normalized_burden_hr'),
 ('Burden Gini','burden_gini'),('Makespan (h)','makespan_hr'),('Total travel (crew-hours)','total_travel_hr')]
table('TABLE2','Table 2a. Baseline means across the same 32 physical realizations.',
 ['Metric']+LAB,[[label]+[f'{mean(s,m):.{4 if m=="burden_gini" else 3}f}' for s in STR] for label,m in metrics],
 'Burden hours are integrated deficits of a modeled service-access proxy, not observed outage hours. See Eq. 2 for the two population denominators.',[2.15,1.15,1.15,1.15,1.2])
table('TABLE2','Table 2b. Paired differences from Hospital-first, with 95% bootstrap intervals.',
 ['Metric','Balanced − HF','HospFirst − HF','Efficiency − HF'],
 [[label]+[effect(s,m,digits=4 if m=='burden_gini' else 3) for s in STR[1:]] for label,m in metrics],
 'Negative means a lower metric than Hospital-first. Intervals resample 32 realization units 10,000 times; no equivalence or simultaneous-testing claim follows.',[1.95,1.62,1.62,1.61])
table('TABLE3','Table 3. Change in the strategy-minus-Hospital-first contrast when DS4 durations are doubled.',
 ['Metric','Balanced','HospFirst','Efficiency'],
 [[label]+[effect(s,m,'DS4x2','change_in_strategy_minus_HF') for s in STR[1:]] for label,m in [metrics[0],('Q4 burden (h)','Q4_burden_hr'),metrics[3],metrics[5],metrics[6]]],
 'Entries are paired differences-in-differences in hours, with 95% bootstrap intervals. Positive indicates that the strategy’s relative cost increases under longer DS4 actions; it is not the total duration-condition effect.',[1.95,1.62,1.62,1.61])
table('STABLE1','Table S1. Tract evidence strata and population.',
 ['Stratum','Candidate composition','Tracts','Population'],[
 ['S1','A only',444,'1,936,737'],['S2','Resolved A/B; no C',318,'1,384,566'],['S3','Resolved A/B with C',43,'199,079'],['S4','C only; normalized burden NA',12,'51,770']],
 'Total: 817 tracts and 3,572,152 people. Classes A/B/C refer to attachment evidence, not utility ownership.',[.65,3.8,.8,1.55])
table('STABLE2','Table S2. Adopted anchored fragility parameters.',
 ['Voltage class (assets)','Parameter','DS1','DS2','DS3','DS4'],[
 ['Low (249)','Median PGA, g','.15','.29','.45','.90'],['Low','Log dispersion β','.70','.55','.45','.45'],
 ['Medium (51)','Median PGA, g','.15','.25','.35','.70'],['Medium','Log dispersion β','.60','.50','.40','.40'],
 ['High (2)','Median PGA, g','.11','.15','.20','.47'],['High','Log dispersion β','.50','.45','.35','.40']],
 'Low: 34.5–<150 kV; medium: 150–<350 kV; high: ≥350 kV. These adopted classes do not establish asset-specific calibration.',[1.7,1.7,.85,.85,.85,.85])
table('STABLE3','Table S3a. Policy-design coefficients.',
 ['Objective','Population','Hospital','SOVI','Makespan'],[
 ['Balanced',1,3,1,.5],['HospFirst',1,20,1,.1],['Efficiency',1,1,1,2]],
 'The first three coefficients normalize service priority; the last penalizes makespan separately. All are design assumptions.',[1.8,1.25,1.25,1.25,1.25])
table('STABLE3','Table S3b. Same-objective candidate scores and provenance.',
 ['Objective','Original returned','Adopted candidate','Hospital-first'],[
 ['Balanced','0.780985337154','0.793690122004','0.792733835572'],
 ['HospFirst','0.900607781997','0.916459946176','0.916319981609'],
 ['Efficiency','0.368272354637','0.368272354637','0.338924957558']],
 'Compare columns within a row only. Balanced/HospFirst adopted candidates are reconstructed deterministic priority initializers; Efficiency is the saved seed-47 search sequence. Unsaved intermediate chromosomes and global optima are not claimed recovered.',[1.3,1.8,1.85,1.85])
table('STABLE4','Table S4. Fixed vulnerability quartiles and metric denominators.',
 ['Quartile','All tracts','All population','Population in R>0 tracts'],[
 ['Q1',205,'812,549','808,424'],['Q2',205,'915,074','907,863'],['Q3',203,'910,124','890,003'],['Q4',204,'934,405','914,092']],
 'Quartiles use tract NRI-derived SOVI_SCORE. Population-normalized group burden uses the final column, not the total population column.',[1.0,1.25,2.0,2.55])
rows=[]
for s,l in zip(STR[1:],LAB[1:]):
    row=[l]
    for stat in ['classification_of_32_realization_mean','single_realization_classification']:
        for cls in ['improved','worsened']:
            x=W[(W.condition=='corrected_baseline')&(W.strategy==s)&(W.group=='ALL')&(W.statistic==stat)&(W.classification==cls)]
            row.append(f'{x.pct_all_group_population.mean():.2f}%')
    rows.append(row)
table('STABLE5','Table S5. Baseline population classification under two averaging operations.',
 ['Sequence','Mean-effect improved','Mean-effect worsened','Realization-first improved','Realization-first worsened'],rows,
 'All percentages use 3,572,152 people. Mean-effect classification applies ±1 h to each tract’s 32-realization mean; realization-first classifies each draw then averages shares. Unresolved is 1.45% in both. The ±1 h threshold is not a significance test.',[1.3,1.35,1.35,1.4,1.4])

FIGS={
 'FIG1':('Figure_1_Study_design','Figure 1. Conditional study design. The same physical realizations feed four fixed sequences. The A/B influence comparison is an offline change of evaluation weights; the DS4 comparison reschedules the same tasks with relatively longer severe-damage actions. Neither adds physical draws or optimizes a new policy.'),
 'FIG2':('Figure_2_Group_burdens','Figure 2. Absolute group burdens and relative distributional measures under baseline assumptions. Q1–Q4 are fixed NRI-derived vulnerability quartiles. Panels B and C show paired mean differences and 95% realization-bootstrap intervals. Panel D shows four evaluated policy points, not a Pareto frontier. Lower Gini for Efficiency coexists with higher burden in every quartile.'),
 'FIG3':('Figure_3_Tract_classification','Figure 3. Population shares classified by burden changes relative to Hospital-first. Panel A classifies each tract’s mean effect over 32 realizations; panel B averages classifications performed separately within each realization. Improved is below −1 h, worsened above +1 h, and near-zero between those thresholds. These are practical classes, not significance tests. Twelve unresolved tracts remain separate; all percentages use the full 3,572,152-person domain.'),
 'FIG4':('Figure_4_Targeted_contrasts','Figure 4. Two separate targeted contrasts. Panels A–B vary Class B influence in mixed A/B tracts, retaining resolved mass and fixed decisions; λ=1 is baseline. Panels C–D double only stored DS4 task durations, with original weights and sequences. The tests are not combined into a factorial experiment. Curves connect evaluated conditions and do not imply interpolation evidence. Paired intervals for the duration contrast appear in Table 3.'),
 'SFIG1':('Figure_S1_Search_records','Figure S1. Retained generation-best fitness records for ten seeds per objective. Dashed lines show the adopted candidate scores. Nonelitist records may decrease; these are not best-so-far curves. Scores cannot be compared across policy objectives. Balanced/HospFirst adopted sequences are reconstructed deterministic initializers; no new search was run.'),
 'SFIG2':('Figure_S2_Denominators_and_coverage','Figure S2. Denominator and representation boundaries. Panel A separates direct-only, resolved A/B, and partly unresolved evidence strata. Panel B compares full-population weighting of normalized tract burdens with population×resolved-mass weighting. These are descriptive strata; fully unresolved tracts have no numeric burden and are not plotted as zero.')}

EQTEXT={
 'EQ1':'Bᵣ = ∫₀ᴴ [Rᵣ − Lᵣ(t)] dt;     Nᵣ = Bᵣ/Rᵣ, for Rᵣ > 0.  (1)',
 'EQ2':'N̄ₚ = ΣᵣPᵣNᵣ / ΣᵣPᵣ;     Bₚᵣ = ΣᵣPᵣBᵣ / ΣᵣPᵣRᵣ.  (2)  Sums are over Rᵣ > 0.',
 'SEQ3':'Fₚ = Σᵢqᵢ max(Tₘₐₓ − Cᵢ, 0) / (Tₘₐₓ Σᵢqᵢ) − wₘ Cₘₐₓ/Tₘₐₓ;  Tₘₐₓ = 504 h.  (S1)'}

def m(tag,*children,text=None):
    e=OxmlElement('m:'+tag)
    if text is not None: e.text=text
    for c in children:e.append(c)
    return e
def mr(t):return m('r',m('t',text=t))
def sub(b,s):return m('sSub',m('e',mr(b)),m('sub',mr(s)))
def frac(a,b):return m('f',m('num',*a),m('den',*b))
def sumterm(*e):return [mr('∑'),*e]
def equation(doc,key):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10)
    om=m('oMath')
    if key=='EQ1':
        integ=m('nary');pr=m('naryPr');ch=m('chr');ch.set(qn('m:val'),'∫');pr.append(ch);integ.append(pr)
        integ.extend([m('sub',mr('0')),m('sup',mr('H')),m('e',mr('['),sub('R','r'),mr(' − '),sub('L','r'),mr('(t)] dt'))])
        nodes=[sub('B','r'),mr(' = '),integ,mr(' ;    '),sub('N','r'),mr(' = '),frac([sub('B','r')],[sub('R','r')]),mr(' ,  '),sub('R','r'),mr(' > 0       (1)')]
    elif key=='EQ2':
        nodes=[sub('N̄','P'),mr(' = '),frac(sumterm(sub('P','r'),sub('N','r')),sumterm(sub('P','r'))),mr(' ;    '),sub('B','PR'),mr(' = '),frac(sumterm(sub('P','r'),sub('B','r')),sumterm(sub('P','r'),sub('R','r'))),mr('       (2)')]
    else:
        nodes=[sub('F','p'),mr(' = '),frac(sumterm(sub('q','i'),mr(' max('),sub('T','max'),mr(' − '),sub('C','i'),mr(', 0)')),[sub('T','max'),*sumterm(sub('q','i'))]),mr(' − '),sub('w','M'),frac([sub('C','max')],[sub('T','max')]),mr('       (S1)')]
    for n in nodes:om.append(n)
    p._p.append(om)
    if key=='EQ2':doc.add_paragraph('All sums in Eq. 2 are over tracts with R > 0.','Note')

review=(ROOT.parent/'01_Reviewer_and_Decision/IJDRR-D-26-02276_editor_review_transfer_record.md').read_text(encoding='utf-8')
COMMENTS={}
for rn,n in [('R1',7),('R2',16)]:
    section=review.split(f'## Reviewer {rn[-1]} — complete report')[1].split('\n## ')[0].split('\n### Transcription')[0]
    paras=re.split(r'\n>\s*\n',section.strip())
    intro=[]
    for para in paras:
        text=' '.join(re.sub(r'^>\s?','',l).strip() for l in para.splitlines()).strip()
        match=re.match(r'^(\d+)\.\s+(.*)',text,re.S)
        if match:COMMENTS[f'{rn}:{match[1]}']=match[2]
        elif text:intro.append(text)
    COMMENTS[f'{rn}:0']='\n\n'.join(intro)
    assert all(f'{rn}:{k}' in COMMENTS for k in range(n+1))

def new_doc(kind):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
    sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
    styles=doc.styles
    for st in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        styles[st].font.name='Times New Roman';styles[st].font.color.rgb=RGBColor(0,0,0)
        rf=styles[st]._element.get_or_add_rPr().find(qn('w:rFonts'))
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            if qn('w:'+attr) in rf.attrib:del rf.attrib[qn('w:'+attr)]
        for border in styles[st]._element.findall('.//'+qn('w:pBdr')):
            border.getparent().remove(border)
    styles['Normal'].font.size=Pt(11)
    styles['Normal'].paragraph_format.line_spacing=1.12
    styles['Normal'].paragraph_format.space_after=Pt(7)
    styles['Title'].font.size=Pt(18);styles['Title'].font.bold=True
    styles['Title'].paragraph_format.space_after=Pt(14)
    for st,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        styles[st].font.size=Pt(size);styles[st].font.bold=True
        styles[st].paragraph_format.space_before=Pt(13);styles[st].paragraph_format.space_after=Pt(6)
    for name,size in [('Note',9.5),('Review quote',10.5)]:
        if name not in styles:styles.add_style(name,1)
        styles[name].base_style=styles['Normal'];styles[name].font.size=Pt(size)
    styles['Review quote'].font.italic=True
    styles['Review quote'].paragraph_format.left_indent=Inches(.15)
    styles['Caption'].font.size=Pt(10);styles['Caption'].font.italic=False;styles['Caption'].font.bold=False
    f=sec.footer.paragraphs[0];f.alignment=2
    f.add_run('Page ').font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    doc.core_properties.title=kind.replace('_',' ').title()
    doc.core_properties.author='Yinchen Yi; Yutong Li; Marta C. González'
    doc.core_properties.subject='Conditional restoration decisions; results based on commit 87110da'
    return doc

def addtable(doc,title,headers,rows,note='',widths=None):
    if title:
        p=doc.add_paragraph(title,'Caption');p.paragraph_format.keep_with_next=True
    tab=doc.add_table(rows=1,cols=len(headers));tab.autofit=False
    if widths is None:widths=[6.8/len(headers)]*len(headers)
    for col,w in zip(tab.columns,widths):col.width=Inches(w)
    for c,h,w in zip(tab.rows[0].cells,headers,widths):c.text=str(h);c.width=Inches(w)
    for row in rows:
        for c,v,w in zip(tab.add_row().cells,row,widths):c.text=str(v);c.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(tab.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for cell in row.cells:
            tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side,val in [('top','65'),('bottom','65'),('left','65'),('right','65')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');mar.append(el)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders');bot=OxmlElement('w:bottom');bot.set(qn('w:val'),'single');bot.set(qn('w:sz'),'4');bot.set(qn('w:color'),'AAAAAA');borders.append(bot);tcpr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
                for run in p.runs:run.font.size=Pt(9.5);run.bold=ri==0
    doc.add_paragraph(note,'Note') if note else doc.add_paragraph()

def mdtable(item):
    title,h,rows,note,_=item
    return '\n\n'+title+'\n\n'+'| '+' | '.join(map(str,h))+' |\n'+'|'+'|'.join(['---']*len(h))+'|\n'+'\n'.join('| '+' | '.join(map(str,r))+' |' for r in rows)+'\n\n'+note+'\n\n'

def build(name):
    raw=(SOURCE/(name+'.md.in')).read_text(encoding='utf-8')
    doc=new_doc(name);complete=[];blocks=re.split(r'\n\s*\n',raw.strip())
    for block in blocks:
        block=block.strip();match=re.fullmatch(r'\[\[(.*?)\]\]',block)
        if match:
            key=match[1]
            if key in TABLES:
                for item in TABLES[key]:addtable(doc,*item);complete.append(mdtable(item))
            elif key in FIGS:
                filename,caption=FIGS[key]
                p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(ROOT/'Figures'/(filename+'.png')),width=Inches(6.7))
                doc.add_paragraph(caption,'Caption')
                complete.append(f'![{caption}](Figures/{filename}.png)\n\n{caption}')
            elif key in EQTEXT:
                equation(doc,key);complete.append(EQTEXT[key])
            elif key.startswith('COMMENT:'):
                text=COMMENTS[key[8:]]
                for para in text.split('\n\n'):doc.add_paragraph(para,'Review quote')
                complete.append('\n\n'.join('> '+p for p in text.split('\n\n')))
            else:raise ValueError(key)
        elif block.startswith('#'):
            heading=re.match(r'^(#+)\s+(.*)',block,re.S);level=len(heading[1]);text=heading[2]
            doc.add_paragraph(text,'Title' if level==1 else f'Heading {level-1}');complete.append(block)
        elif block.startswith('|'):
            lines=block.splitlines();head=[x.strip() for x in lines[0].strip('|').split('|')]
            rows=[[x.strip() for x in l.strip('|').split('|')] for l in lines[2:]]
            addtable(doc,'',head,rows,widths=[1.8,2.25,2.75]);complete.append(block)
        else:
            p=doc.add_paragraph(block.replace('\n',' '))
            if block.startswith('Response:'):
                p.runs[0].text=block[len('Response:'):].strip();r=OxmlElement('w:r');pr=OxmlElement('w:rPr');pr.append(OxmlElement('w:b'));r.append(pr);t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text='Response: ';r.append(t);p._p.insert(0,r)
            if re.match(r'^[A-ZÇa-z].*\(\d{4}\).*https://',block):p.paragraph_format.space_after=Pt(6)
            complete.append(block)
    if name=='CLAIM_EVIDENCE_CROSSWALK':
        for tab in doc.tables:
            for row in tab.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0)
                        for r in p.runs:r.font.size=Pt(9)
                    for el in cell._tc.findall('.//'+qn('w:tcMar')+'/*'):
                        if el.tag in [qn('w:top'),qn('w:bottom')]:el.set(qn('w:w'),'35')
    doc.save(ROOT/(name+'.docx'))
    finished='\n\n'.join(complete)+'\n';assert '[[' not in finished
    (ROOT/(name+'.md')).write_text(finished,encoding='utf-8')
    print(name, 'paragraphs',len(doc.paragraphs),'tables',len(doc.tables))

for name in NAMES:build(name)
stats=ROOT/'Supplementary_tables';stats.mkdir(exist_ok=True)
for f in ['METRIC_DISTRIBUTIONS.csv','PAIRED_EFFECTS.csv','METRIC_RANK_FREQUENCIES.csv']:
    if (SRC/f).exists():shutil.copyfile(SRC/f,stats/f)
print('Original reviewer comments embedded:',len(COMMENTS))
