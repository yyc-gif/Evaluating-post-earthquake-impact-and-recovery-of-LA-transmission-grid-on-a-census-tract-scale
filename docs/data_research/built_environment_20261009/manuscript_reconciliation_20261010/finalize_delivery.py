"""Check the review artifacts and freeze a package manifest after visual review."""
import argparse
import csv
import hashlib
import json
import os
from pathlib import Path

import docx
import pandas as pd
from pypdf import PdfReader


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(path, value):
    path.write_text(json.dumps(value, indent=2) + "\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", type=Path, required=True)
    a = ap.parse_args()
    repo = a.repo.resolve()
    if os.name == "nt":
        repo = Path("\\\\?\\" + str(repo))
    base = repo / "docs/data_research/built_environment_20261009"
    out = base / "manuscript_reconciliation_20261010"
    source = base / "manuscript_integration_preview_20261010"
    numeric = json.loads((out / "NUMERICAL_QA.json").read_text())
    structure = json.loads((out / "DOCUMENT_STRUCTURE_QA.json").read_text())
    render = json.loads((out / "DOCUMENT_RENDER_QA.json").read_text())
    preservation = json.loads((out / "PRESERVATION_AFTER.json").read_text())
    assert numeric["status"] == preservation["status"] == "PASS"
    assert preservation["protected_files"] == 950 and preservation["staged_item_count"] == 862
    assert render["visual_review_confirmed"] and render["documents"]["review"]["page_count"] == 20
    assert len(structure["original_paragraphs_preserved"]) == 108
    assert structure["original_tables_preserved"] == [1, 2, 3]
    review_path = out / "Manuscript_Stage7_Integration_REVIEW.docx"
    assert sha(review_path) == structure["review_sha256"]
    pdf_path = out / "Manuscript_Stage7_Integration_REVIEW.pdf"
    assert sha(pdf_path) == render["documents"]["review"]["pdf_sha256"]
    d = docx.Document(str(review_path))
    profiles = pd.read_csv(source / "STAGE7_CLUSTER_PROFILES.csv").set_index("candidate_cluster")
    table_checks = []
    fields = [["n_tracts", "population_sum", "B_480_hr_mean", "T80_mean", "Init_Supply_mean"],
              ["SOVI_SCORE_mean", "land_use_entropy_mean", "building_footprint_coverage_mean",
               "all_use2014_pre1970_area_share_mean", "NRI_BUILDVALUE_median", "ALR_NPCTL_median"]]
    for table, columns in zip(d.tables[3:], fields):
        for row, cluster in zip(table.rows[1:], ["C1", "C2", "C3", "C4"]):
            assert row.cells[0].text == cluster
            for cell, field in zip(row.cells[1:], columns):
                v = profiles.loc[cluster, field]
                if field in ["n_tracts", "population_sum"]:
                    expected = format(int(v), ",")
                elif field == "Init_Supply_mean":
                    expected = format(v * 100, ".5f")
                elif field == "NRI_BUILDVALUE_median":
                    expected = format(v / 1e6, ".1f")
                elif "fraction" in field or "coverage" in field or "entropy" in field or "area_share" in field:
                    expected = format(v, ".3f")
                else:
                    expected = format(v, ".2f")
                assert cell.text == expected, (cluster, field, cell.text, expected)
                table_checks.append({"class": cluster, "field": field, "display": expected, "pass": True})
    text = "\n".join(p.text for p in d.paragraphs)
    for phrase in ["No PCA reduction defines this partition", "0.216624", "36.96%",
                   "850", "2,287", "12,390", "1.815%", "sqrt(0.20/p)", "98.1563%"]:
        assert phrase in text, phrase
    for phrase in ["The retained five", "The retained Stage 7 PCA/K-means", "36(3), 265"]:
        assert phrase not in text, phrase
    for number in range(1, 9):
        assert sum(p.text.startswith("Figure " + str(number) + ".") for p in d.paragraphs) == 1
    assert "Figure S_STAGE7" in text and len(d.tables) == 5
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(out / "FIG_STAGE7_TYPOLOGY_MAP_REVIEW.pdf").pages)
    assert "Outside cohort (24 tracts)" in pdf_text and "No residents (24 tracts)" not in pdf_text
    structure["actual_rendering"] = "PASS: native Word export, 20 review pages visually inspected"
    dump(out / "DOCUMENT_STRUCTURE_QA.json", structure)
    figures = json.loads((out / "FIGURE_RENDER_QA.json").read_text())
    for record in figures:
        record["manual_inspection"] = "PASS"
        record["notes"] = "Actual 185mm/178mm pages inspected; no clipped labels/legends. Review Word figures also checked at172.72mm. Original map legend has a documented review-only correction."
    dump(out / "FIGURE_RENDER_QA.json", figures)
    dump(out / "DELIVERY_QA.json", {"status": "PASS_FOR_AUTHOR_REVIEW", "formal_promotion": False,
         "submitted": False, "numerical_checks": len(numeric["checks"]), "rounded_new_table_checks": table_checks,
         "review_pages_inspected": 20, "original_paragraphs_preserved": 108, "original_tables_preserved": 3,
         "original_figure_payloads_preserved": 6, "protected_science_files": 950, "original_staged_items": 862,
         "journal_confirmed": False, "supplementary_final_numbering": False,
         "unresolved": ["Separate hotspot revalidation", "September/October GA experimental scope reconciliation",
                        "Final journal/supplementary numbering", "Partial Lloyd publisher verification",
                        "Archived SCAG version conflict and LARIAC completeness caveats"]})
    excluded = {"OUTPUT_HASHES.csv", "REMOTE_VERIFICATION.json"}
    rows = []
    for path in sorted(out.rglob("*")):
        if not path.is_file() or path.name in excluded or "rendered_qa" in path.parts or "__pycache__" in path.parts or path.suffix == ".log":
            continue
        rows.append({"file": path.relative_to(out).as_posix(), "bytes": path.stat().st_size, "sha256": sha(path)})
    with (out / "OUTPUT_HASHES.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["file", "bytes", "sha256"])
        writer.writeheader()
        writer.writerows(rows)
    print("PASS author-review QA;", len(table_checks), "rounded table values;", len(rows), "manifested payloads")


if __name__ == "__main__":
    main()
